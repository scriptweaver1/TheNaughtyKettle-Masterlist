#!/usr/bin/env python3
"""
Build the site's audios.json from the masterlist spreadsheet.

    python3 build_audios_json.py theNaughtyKettle_Audios_MERGED_fixed.xlsx
    python3 build_audios_json.py sheet.xlsx -o site/audios.json

Re-run this whenever the spreadsheet changes; it overwrites audios.json.
Entry ids are slugs of the title, so they stay stable across rebuilds and
visitors' saved favourites keep working.

Spreadsheet column  ->  JSON field
    Title               title   (an "[EXCLUSIVE]" prefix is stripped and
                                 becomes exclusive: true instead)
    Tags                tags[]
    Date                date        Type            type
    Writer              writer      Description     description
    Editor              editor      Duration        duration
    Collab Partners     collab[]
    Post Link / Script Link / Soundgasm / Audiochan /
    Hot Audio / Patreon / C4S   ->  links{}

"x" is treated as "nothing here" in every column, matching the sheet's
own convention, and becomes null/omitted rather than the literal string.
"""

import argparse, json, re, sys, unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime

import openpyxl

# Cells that mean "nothing here" rather than a value.
BLANK = {"", "x", "-", "--", "n/a", "na", "none", "#n/a", "?", "tbd"}

# Spreadsheet column -> key used in the JSON "links" object. Order decides
# the order the buttons appear on a card.
LINK_COLUMNS = [
    ("Post Link", "reddit"),
    ("Soundgasm", "soundgasm"),
    ("Audiochan", "audiochan"),
    ("Hot Audio", "hotaudio"),
    ("Patreon", "patreon"),
    ("C4S", "c4s"),
    ("Script Link", "script"),
]

EXCLUSIVE_RE = re.compile(r"^\s*[\[\(]\s*exclusive\s*[\]\)]\s*[-–:]?\s*", re.I)
# "Oil My Pussy 6: ...", "Audio Postcards 2- ...", "Kettlecast #1: ..."
SERIES_RE = re.compile(r"^(.{3,44}?)\s*#?\s*(\d+(?:\.\d+)?)\s*[:\-–]\s*\S")


def clean(value):
    """Normalise a cell to a string, with the sheet's placeholders as ''."""
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    text = str(value).strip()
    return "" if text.lower() in BLANK else text


def to_iso_date(value):
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    text = clean(value)
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%m/%d/%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    return text  # unparseable: pass through rather than silently dropping


def slugify(text, fallback="entry"):
    text = unicodedata.normalize("NFKD", text)
    text = "".join(c for c in text if not unicodedata.category(c).startswith("So"))
    text = re.sub(r"[^a-zA-Z0-9]+", "-", text).strip("-").lower()
    return text[:60].strip("-") or fallback


def split_people(text):
    """'A, B' / 'A & B' / 'A x B' -> ['A', 'B'], placeholders dropped."""
    if not text:
        return []
    parts = re.split(r"[,/;&]| \bx\b | \+ ", text)
    out = []
    for part in parts:
        name = part.strip().lstrip("@").strip()
        name = re.sub(r"^(?:/?u/)", "", name, flags=re.I).strip()
        if name and name.lower() not in BLANK and name not in out:
            out.append(name)
    return out


def person(name):
    """A credit, linked to the Reddit profile the username implies."""
    return {"name": name, "url": f"https://www.reddit.com/user/{name}"}


def split_tags(text):
    out = []
    seen = set()
    for raw in re.split(r"[,•|]", text or ""):
        tag = raw.strip().strip("[]()").strip()
        if not tag or tag.lower() in BLANK:
            continue
        if tag.lower() not in seen:
            seen.add(tag.lower())
            out.append(tag)
    return out


def normalise_url(url):
    if not url:
        return None
    url = url.strip()
    if url.startswith("//"):
        return "https:" + url
    if not re.match(r"^https?://", url, re.I):
        if re.match(r"^(www\.|reddit\.com|soundgasm|audiochan|hotaudio|patreon|clips4sale)", url, re.I):
            return "https://" + url
        return None
    return re.sub(r"[?&]utm_[^&]*", "", url).rstrip("?&")


def read_rows(path, sheet=None):
    book = openpyxl.load_workbook(path, data_only=True)
    page = book[sheet] if sheet else book[book.sheetnames[0]]
    header = [c.value for c in page[1]]
    rows = []
    for values in page.iter_rows(min_row=2, values_only=True):
        if all(v in (None, "") for v in values):
            continue
        rows.append({(h or f"col{i}"): v for i, (h, v) in enumerate(zip(header, values))})
    return rows


def detect_series(titles):
    """Map title -> series name, but only where a prefix is shared by 2+ entries."""
    groups = defaultdict(list)
    for title in titles:
        match = SERIES_RE.match(title)
        if match:
            name = match.group(1).strip(" -–:#")
            # "Audio Postcard" and "Audio Postcards" are the same series
            groups[re.sub(r"s$", "", name.lower())].append((title, name))
    series = {}
    for members in groups.values():
        if len(members) < 2:
            continue
        # use the most common spelling of the name
        label = Counter(name for _, name in members).most_common(1)[0][0]
        for title, _ in members:
            series[title] = label
    return series


def build(rows):
    titles = [clean(r.get("Title")) for r in rows]
    series_of = detect_series([t for t in titles if t])

    entries, used_ids, notes = [], set(), Counter()

    for row in rows:
        raw_title = clean(row.get("Title"))
        if not raw_title:
            notes["skipped: no title"] += 1
            continue

        exclusive = bool(EXCLUSIVE_RE.match(raw_title))
        title = EXCLUSIVE_RE.sub("", raw_title).strip() if exclusive else raw_title

        entry_id = slugify(title)
        if entry_id in used_ids:                       # two versions of one title
            n = 2
            while f"{entry_id}-{n}" in used_ids:
                n += 1
            entry_id = f"{entry_id}-{n}"
            notes["duplicate title, id suffixed"] += 1
        used_ids.add(entry_id)

        links = {}
        for column, key in LINK_COLUMNS:
            url = normalise_url(clean(row.get(column)))
            if url:
                links[key] = url
        if not links:
            notes["no links at all"] += 1

        writers = split_people(clean(row.get("Writer")))
        editors = split_people(clean(row.get("Editor")))
        collab = split_people(clean(row.get("Collab Partners")))
        if len(writers) > 1:
            notes["multiple writers (joined)"] += 1
        if len(editors) > 1:
            notes["multiple editors (joined)"] += 1

        def credit(names):
            if not names:
                return None
            if len(names) == 1:
                return person(names[0])
            # the card renders one credit line, so join and link the first
            return {"name": ", ".join(names), "url": person(names[0])["url"]}

        entry_date = to_iso_date(row.get("Date"))
        if not entry_date:
            notes["no date"] += 1

        duration = clean(row.get("Duration"))
        description = clean(row.get("Description"))
        if not description:
            notes["no description"] += 1

        entries.append({
            "id": entry_id,
            "title": title,
            "date": entry_date,
            "duration": duration or None,
            "series": series_of.get(raw_title),
            "type": clean(row.get("Type")) or "Uncategorized",
            "tags": split_tags(clean(row.get("Tags"))),
            "description": description or None,
            "image": None,
            "exclusive": exclusive,
            "writer": credit(writers),
            "editor": credit(editors),
            "collab": [person(n) for n in collab] or None,
            "links": links,
        })

    entries.sort(key=lambda e: (e["date"] or "0000-00-00"), reverse=True)
    return entries, notes


def main():
    ap = argparse.ArgumentParser(description="Build audios.json from the masterlist spreadsheet.")
    ap.add_argument("xlsx", help="masterlist .xlsx")
    ap.add_argument("-o", "--out", default="audios.json", help="output path (default: audios.json)")
    ap.add_argument("--sheet", default=None, help="sheet name (default: the first one)")
    ap.add_argument("--indent", type=int, default=1, help="JSON indent; 0 for a minified file")
    args = ap.parse_args()

    rows = read_rows(args.xlsx, args.sheet)
    entries, notes = build(rows)

    with open(args.out, "w", encoding="utf-8") as fh:
        json.dump(entries, fh, ensure_ascii=False,
                  indent=args.indent or None,
                  separators=(",", ":") if not args.indent else None)

    types = Counter(e["type"] for e in entries)
    linked = Counter(k for e in entries for k in e["links"])

    print(f"read  {len(rows)} spreadsheet rows")
    print(f"wrote {len(entries)} entries -> {args.out}")
    if entries:
        print(f"dates {entries[-1]['date']} .. {entries[0]['date']}")
    print("types:      " + ", ".join(f"{k} {v}" for k, v in types.most_common()))
    print("links:      " + ", ".join(f"{k} {v}" for k, v in linked.most_common()))
    print(f"series:     {sum(1 for e in entries if e['series'])} entries in "
          f"{len({e['series'] for e in entries if e['series']})} series")
    print(f"exclusive:  {sum(1 for e in entries if e['exclusive'])}")
    print(f"credits:    writer {sum(1 for e in entries if e['writer'])}, "
          f"editor {sum(1 for e in entries if e['editor'])}, "
          f"collab {sum(1 for e in entries if e['collab'])}")
    if notes:
        print("worth a look:")
        for label, count in notes.most_common():
            print(f"    {count:4d}  {label}")


if __name__ == "__main__":
    sys.exit(main())
